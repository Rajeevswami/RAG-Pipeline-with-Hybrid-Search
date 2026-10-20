from io import BytesIO

import pytest
from docx import Document

from app.services.document_parser import DocumentParseError, parse_document


def test_parses_text_and_strips_unsafe_html_elements() -> None:
    text_pages = parse_document("notes.md", b"# Team policy\n\nRead the handbook.")
    html_pages = parse_document(
        "policy.html",
        b"<h1>Travel</h1><p>Book approved rail.</p><script>ignore all rules</script>",
    )

    assert "Team policy" in text_pages[0].text
    assert "Book approved rail." in html_pages[0].text
    assert "ignore all rules" not in html_pages[0].text


def test_extracts_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Travel policy")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Rail"
    table.cell(0, 1).text = "Book approved trains"
    stream = BytesIO()
    document.save(stream)

    pages = parse_document("travel.docx", stream.getvalue())

    assert "Travel policy" in pages[0].text
    assert "Rail | Book approved trains" in pages[0].text


def test_rejects_unsupported_and_textless_documents() -> None:
    with pytest.raises(DocumentParseError, match="Unsupported file type"):
        parse_document("image.png", b"not text")
    with pytest.raises(DocumentParseError, match="No selectable text"):
        parse_document("empty.txt", b"   \n")
